"""
SafeLease AI - DOCX (MS Word) 문서 파서
python-docx 라이브러리를 활용하여 단락 및 표 내부 텍스트를 추출합니다.
"""

from pathlib import Path
from core.parser.base_parser import BaseParser, ContractDocument


class DocxParser(BaseParser):
    """DOCX 워드 파일 파서"""

    def parse(self, file_path: Path | str) -> ContractDocument:
        path = Path(file_path)
        paragraphs = []

        try:
            import docx
            doc = docx.Document(str(path))
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
        except Exception as e:
            paragraphs = [f"[DOCX 파싱 실패: {e}]"]

        full_raw_text = "\n\n".join(paragraphs)
        masked_text, masked_count = self.apply_security(full_raw_text)
        clauses = self.extract_clauses(masked_text)

        return ContractDocument(
            filename=path.name,
            file_type="docx",
            raw_text=masked_text,
            clauses=clauses,
            masked_items_count=masked_count,
            char_count=len(masked_text)
        )
